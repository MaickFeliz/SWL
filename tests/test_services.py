"""Pruebas unitarias de la capa de servicios.

Cubre:
- ReportService.generate_overdue_users_report() con base vacía.
- LoanService.approve_loan() validando transición de estados Enum.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

import openpyxl
import pytest
from flask import Flask

from app import db
from app.models import Catalog, ItemInstance, InventoryStatus, Loan, LoanStatus, User
from app.services.loan_service import LoanService
from app.services.report_service import ReportService


# ---------------------------------------------------------------------------
# ReportService
# ---------------------------------------------------------------------------

class TestGenerateOverdueUsersReport:
    """Valida el comportamiento de ReportService.generate_overdue_users_report()."""

    def test_returns_report_with_no_overdue_rows(self, app_context: Flask) -> None:
        """Debe retornar un BytesIO con un Excel válido y sin filas de morosos cuando
        no existe ningún préstamo con estado OVERDUE en la base de datos.

        Garantiza que el reporte no lanza excepciones ante tablas vacías o sin
        morosos, y que el archivo resultante puede abrirse con openpyxl sin errores.
        """
        output: BytesIO = ReportService.generate_overdue_users_report()

        assert isinstance(output, BytesIO), "El resultado debe ser un objeto BytesIO"
        assert output.tell() == 0, "El puntero debe estar al inicio (seek(0) aplicado)"

        # Verificar que es un Excel válido y parseable
        workbook = openpyxl.load_workbook(output)
        assert "Usuarios Morosos" in workbook.sheetnames

        sheet = workbook["Usuarios Morosos"]
        # Solo debe existir la fila de encabezados (o estar completamente vacía)
        data_rows = [
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(cell is not None for cell in row)
        ]
        assert len(data_rows) == 0, (
            "No deben existir filas de datos cuando la base está vacía"
        )


# ---------------------------------------------------------------------------
# LoanService
# ---------------------------------------------------------------------------

class TestApproveLoan:
    """Valida la lógica de aprobación de préstamos en LoanService."""

    def _create_user(self) -> User:
        user = User(
            document_id="TEST001",
            full_name="Usuario Prueba",
            role="cliente",
            password_hash="hash_no_importa",
        )
        db.session.add(user)
        db.session.flush()
        return user

    def _create_catalog_and_instance(self) -> ItemInstance:
        catalog = Catalog(
            title_or_name="Laptop Test",
            category="computo",
        )
        db.session.add(catalog)
        db.session.flush()

        instance = ItemInstance(
            catalog_id=catalog.id,
            unique_code="UNIT-TEST-001",
            status=InventoryStatus.LOANED,
        )
        db.session.add(instance)
        db.session.flush()
        return instance

    def test_approve_loan_transitions_status_from_pending_to_active(
        self, app_context: Flask
    ) -> None:
        """Aprobar un préstamo PENDING debe cambiar su estado a ACTIVE.

        Verifica la transición de Enum LoanStatus.PENDING → LoanStatus.ACTIVE
        y que se registre la fecha de aprobación.
        """
        user = self._create_user()
        instance = self._create_catalog_and_instance()

        loan = Loan(
            user_id=user.id,
            instance_id=instance.id,
            status=LoanStatus.PENDING,
            due_date=datetime.now(timezone.utc) + timedelta(days=15),
        )
        db.session.add(loan)
        db.session.commit()

        success, message = LoanService.approve_loan(loan.id)

        assert success is True, f"La aprobación falló: {message}"
        assert loan.status is LoanStatus.ACTIVE, (
            f"Se esperaba LoanStatus.ACTIVE, se obtuvo {loan.status}"
        )
        assert loan.approval_date is not None, (
            "La fecha de aprobación debe registrarse al aprobar el préstamo"
        )

    def test_approve_loan_fails_when_already_active(
        self, app_context: Flask
    ) -> None:
        """Intentar aprobar un préstamo ACTIVE debe retornar fallo sin modificar el estado."""
        user = User(
            document_id="TEST002",
            full_name="Usuario Prueba 2",
            role="cliente",
            password_hash="hash",
        )
        db.session.add(user)
        db.session.flush()

        catalog = Catalog(title_or_name="Libro Test", category="libro")
        db.session.add(catalog)
        db.session.flush()

        instance = ItemInstance(
            catalog_id=catalog.id,
            unique_code="UNIT-TEST-002",
            status=InventoryStatus.LOANED,
        )
        db.session.add(instance)
        db.session.flush()

        loan = Loan(
            user_id=user.id,
            instance_id=instance.id,
            status=LoanStatus.ACTIVE,
            due_date=datetime.now(timezone.utc) + timedelta(days=7),
        )
        db.session.add(loan)
        db.session.commit()

        success, _ = LoanService.approve_loan(loan.id)

        assert success is False, "No debe poderse aprobar un préstamo que ya está ACTIVE"
        assert loan.status is LoanStatus.ACTIVE, "El estado no debe cambiar"
